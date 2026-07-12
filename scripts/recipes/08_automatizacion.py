CELLS = [
{"type": "markdown", "source": r"""
# 08 · Automatización y producción

El último paso del trabajo de un analista senior no es "el notebook corrió
bien una vez" — es que el pipeline **se pueda correr de forma confiable,
repetida y desatendida**. Este módulo usa un ETL real del proyecto
(`scripts/etl_credito.py`) para mostrar el ciclo completo: validación de
esquema, testing, reportes automáticos, configuración segura y scheduling.

## Contenido
1. Anatomía de un pipeline ETL modular y testeable
2. Validación de esquema con `pandera` (fallar rápido y con mensajes claros)
3. Testing de pipelines de datos con `pytest`
4. Reportes automáticos (Excel formateado)
5. Configuración y secretos: nunca hardcodear credenciales
6. Scheduling: cron, APScheduler, Airflow
7. Empaquetado y reproducibilidad: entornos, Docker, CLI
8. Buenas prácticas de Git para proyectos de datos
"""},

{"type": "markdown", "source": r"""
## 1. Anatomía de un pipeline ETL modular

`scripts/etl_credito.py` (en este mismo repo) separa el pipeline en
funciones puras: `extract`, `validar_esquema`, `transform`, `load`. Esto no
es dogma académico — es lo que hace posible **testear cada paso por
separado** y **reusar** la lógica tanto desde un cron job como desde un
notebook de exploración.

```python
def ejecutar_pipeline(path_entrada: Path, path_salida: Path) -> pd.DataFrame:
    df_crudo = extract(path_entrada)
    df_validado = validar_esquema(df_crudo)
    df_transformado = transform(df_validado)
    load(df_transformado, path_salida)
    return df_transformado
```

El antipatrón común es un notebook de 200 celdas donde todo está mezclado:
imposible de testear, de reusar o de programar como job.
"""},

{"type": "code", "source": r"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path("scripts").resolve()))
import logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)-8s | %(message)s", datefmt="%H:%M:%S")

from etl_credito import extract, transform, load, validar_esquema, ejecutar_pipeline

df_resultado = ejecutar_pipeline(Path("data/credito_solicitudes.csv"), Path("data/_output/credito_procesado.csv"))
df_resultado.shape
"""},

{"type": "markdown", "source": r"""
## 2. Validación de esquema con `pandera`

Sin un esquema explícito, un cambio silencioso en la fuente de datos (una
columna que empieza a llegar con nulos, un score fuera de rango) se detecta
tarde — cuando ya rompió un reporte o un modelo. `pandera` define el
contrato de datos como código y falla ruidosamente si no se cumple.
"""},

{"type": "code", "source": r"""
import pandas as pd

# Simulamos un problema de calidad de datos real: un score corrupto (típico de un bug de ingesta)
df_con_error = df_resultado.copy()
df_con_error.loc[0, "buro_score"] = 1500  # fuera de rango válido (300-850)

try:
    validar_esquema(df_con_error)
except Exception as e:
    print(f"Validación falló como se esperaba:\n{type(e).__name__}: {str(e)[:300]}")
"""},

{"type": "markdown", "source": r"""
## 3. Testing de pipelines de datos con `pytest`

`scripts/test_etl_credito.py` testea **invariantes del resultado**, no solo
"que no truene": sin duplicados, sin nulos donde no deberían existir, rangos
válidos. Se corre igual que cualquier suite de tests de software.
"""},

{"type": "code", "source": r"""
import subprocess

resultado = subprocess.run(
    [sys.executable, "-m", "pytest", "scripts/test_etl_credito.py", "-v", "--no-header"],
    capture_output=True, text=True, cwd=Path.cwd(),
)
print(resultado.stdout[-1500:])
if resultado.returncode != 0:
    print(resultado.stderr[-1000:])
"""},

{"type": "markdown", "source": r"""
> En CI/CD (GitHub Actions, GitLab CI), estos tests corren automáticamente en
> cada push — un pipeline que rompe un test nunca llega a producción.
> ```yaml
> # .github/workflows/tests.yml
> on: [push]
> jobs:
>   test:
>     runs-on: ubuntu-latest
>     steps:
>       - uses: actions/checkout@v4
>       - run: pip install -r requirements.txt
>       - run: pytest scripts/ -v
> ```
"""},

{"type": "markdown", "source": r"""
## 4. Reportes automáticos (Excel formateado)

Un reporte que un stakeholder abre cada semana no debería requerir que
alguien lo arme a mano. `openpyxl` permite generar Excel con formato
condicional, formato de número y múltiples hojas — 100% scriptable.
"""},

{"type": "code", "source": r"""
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

resumen_sector = df_resultado.groupby("sector", observed=True).agg(
    n_solicitudes=("solicitud_id", "count"),
    monto_promedio=("monto_solicitado", "mean"),
    tasa_default=("default_90d", "mean"),
).round(2).sort_values("tasa_default", ascending=False).reset_index()

path_reporte = Path("data/_output/reporte_semanal.xlsx")
path_reporte.parent.mkdir(parents=True, exist_ok=True)

with pd.ExcelWriter(path_reporte, engine="openpyxl") as writer:
    resumen_sector.to_excel(writer, sheet_name="Resumen por sector", index=False)
    ws = writer.sheets["Resumen por sector"]

    # Encabezado con formato
    for col_idx, col_name in enumerate(resumen_sector.columns, start=1):
        celda = ws.cell(row=1, column=col_idx)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2A78D6")
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    # Formato condicional simple: resaltar sectores con tasa_default > 20%
    col_tasa = resumen_sector.columns.get_loc("tasa_default") + 1
    for row_idx in range(2, len(resumen_sector) + 2):
        if ws.cell(row=row_idx, column=col_tasa).value > 0.20:
            ws.cell(row=row_idx, column=col_tasa).fill = PatternFill("solid", fgColor="FADBD8")

print(f"Reporte generado: {path_reporte}")
print(resumen_sector)
"""},

{"type": "markdown", "source": r"""
Para distribuir el reporte automáticamente por correo (ej. cada lunes a las
8am vía cron):

```python
import smtplib
from email.message import EmailMessage

def enviar_reporte(destinatarios: list[str], path_adjunto: Path):
    msg = EmailMessage()
    msg["Subject"] = "Reporte semanal — Solicitudes de crédito"
    msg["From"] = "reportes@empresa.com"
    msg["To"] = ", ".join(destinatarios)
    msg.set_content("Reporte automático adjunto. Generado por el pipeline de datos.")

    with open(path_adjunto, "rb") as f:
        msg.add_attachment(
            f.read(), maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=path_adjunto.name,
        )

    # Credenciales SIEMPRE desde variables de entorno, nunca hardcodeadas (ver sección 5)
    with smtplib.SMTP_SSL("smtp.empresa.com", 465) as smtp:
        smtp.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
        smtp.send_message(msg)
```
"""},

{"type": "markdown", "source": r"""
## 5. Configuración y secretos

**Nunca** hardcodear contraseñas, tokens o connection strings en el código.
Patrón estándar: variables de entorno + `.env` para desarrollo local
(el archivo `.env` va en `.gitignore`, nunca se commitea).
"""},

{"type": "code", "source": r"""
import os
from dotenv import load_dotenv

# .env (NO se commitea a git):
#   DB_HOST=localhost
#   DB_USER=analista
#   DB_PASSWORD=***
#   SMTP_USER=reportes@empresa.com

load_dotenv()  # carga variables de .env al entorno si el archivo existe; en prod ya vienen del sistema

db_host = os.environ.get("DB_HOST", "localhost")  # default sensato si no está configurado
print(f"DB_HOST configurado: {db_host}")
print("\nPatrón de acceso seguro a credenciales:")
print('  password = os.environ["DB_PASSWORD"]  # KeyError explícito si falta, mejor que fallar en silencio')
"""},

{"type": "markdown", "source": r"""
## 6. Scheduling: cron, APScheduler, Airflow

- **cron** (Linux/Mac, o Task Scheduler en Windows): suficiente para jobs
  simples de un solo script.
  ```bash
  # Correr el ETL todos los días a las 6:00 AM
  0 6 * * * cd /ruta/proyecto && /ruta/venv/bin/python scripts/etl_credito.py --input data/in.csv --output data/out.csv >> logs/etl.log 2>&1
  ```
- **APScheduler**: scheduling dentro de un proceso Python de larga duración
  (útil si el scheduler vive junto a una app).
- **Airflow / Dagster / Prefect**: cuando hay múltiples pipelines con
  dependencias entre sí, reintentos, alertas y necesitas un dashboard de
  monitoreo — el estándar de facto en equipos de datos maduros.
"""},

{"type": "code", "source": r"""
from apscheduler.schedulers.background import BackgroundScheduler
from apscheduler.triggers.cron import CronTrigger
import time

def job_etl_diario():
    logging.getLogger("scheduler").info("Ejecutando ETL programado...")

scheduler = BackgroundScheduler()
# Trigger real de cron: todos los días a las 6:00 AM (aquí solo se registra, no se espera)
scheduler.add_job(job_etl_diario, CronTrigger(hour=6, minute=0), id="etl_diario")
scheduler.start()

print("Job programado:", scheduler.get_jobs())
scheduler.shutdown(wait=False)  # apagamos inmediatamente: esto es solo demostrativo en el notebook
"""},

{"type": "markdown", "source": r"""
Referencia de un DAG de Airflow equivalente (no se ejecuta aquí — Airflow
corre como servicio):

```python
from airflow import DAG
from airflow.operators.python import PythonOperator
from datetime import datetime

with DAG("etl_credito_diario", start_date=datetime(2024, 1, 1), schedule="0 6 * * *", catchup=False) as dag:
    extraer = PythonOperator(task_id="extract", python_callable=extract)
    transformar = PythonOperator(task_id="transform", python_callable=transform)
    cargar = PythonOperator(task_id="load", python_callable=load)

    extraer >> transformar >> cargar  # dependencias explícitas entre tareas
```
"""},

{"type": "markdown", "source": r"""
## 7. Empaquetado y reproducibilidad

- **`requirements.txt`** (o mejor, `pyproject.toml` con Poetry/uv): fija
  versiones exactas (`pandas==2.2.0`, no `pandas`) para que "funciona en mi
  máquina" no sea un problema en producción.
- **Entorno virtual** (`venv`, `conda`, `uv venv`): aislar dependencias por
  proyecto — nunca instalar librerías de análisis en el Python del sistema.
- **Docker**: cuando el pipeline debe correr igual en cualquier máquina
  (laptop, CI, servidor), sin depender de qué esté instalado en el host.

```dockerfile
FROM python:3.11-slim
WORKDIR /app
COPY requirements.txt .
RUN pip install --no-cache-dir -r requirements.txt
COPY . .
CMD ["python", "scripts/etl_credito.py", "--input", "data/in.csv", "--output", "data/out.csv"]
```

Y una CLI con `click` (usada en `scripts/etl_credito.py`) para que el script
se pueda invocar con parámetros en vez de tener rutas hardcodeadas:

```python
import click

@click.command()
@click.option("--input", "path_entrada", required=True, type=click.Path(exists=True))
@click.option("--output", "path_salida", required=True)
def cli(path_entrada, path_salida):
    ejecutar_pipeline(Path(path_entrada), Path(path_salida))

# uso: python scripts/etl_credito.py --input data/in.csv --output data/out.csv
```
"""},

{"type": "markdown", "source": r"""
## 8. Buenas prácticas de Git para proyectos de datos

- **`.gitignore`** para datos pesados y secretos: nunca commitear `.env`,
  credenciales, ni CSVs de más de unos MB (usar DVC o Git LFS si el
  versionado de datos es un requisito real).
  ```
  .env
  *.csv
  data/_output/
  __pycache__/
  .ipynb_checkpoints/
  *.pkl
  ```
- **Notebooks en control de versiones**: limpiar outputs pesados antes de
  commitear (`nbstripout` o `jupyter nbconvert --clear-output`) — si no, cada
  commit pesa MBs por imágenes embebidas y los diffs son ilegibles.
- **Mensajes de commit** que expliquen el *por qué*, no el *qué* (el diff ya
  muestra el qué): `"Excluir outliers de monto >P99 antes de entrenar: sesgaban el modelo hacia clientes corporativos"` es mejor que `"fix bug"`.
- **Un notebook no es el artefacto final** de un pipeline de producción — es
  para explorar. La lógica que se repite (como este ETL) se gradúa a un
  módulo `.py` testeable, que el notebook simplemente importa y usa.
"""},

{"type": "markdown", "source": r"""
## Resumen y cierre del curso

Con los 8 módulos completos, el flujo de un analista de datos senior queda
cubierto de punta a punta:

`Fundamentos de Python` → `Librerías (NumPy/Pandas)` → `SQL` →
`Estadística aplicada` → `EDA` → `Visualización` → `Machine Learning` →
`Automatización y producción`

El patrón que se repite en los 8 módulos: **nunca confíes en que los datos
están limpios, nunca reportes una métrica sin comparar contra un baseline, y
nunca dejes en un notebook lo que debería vivir en un script testeado.**
"""},
]
